import http.server
import json
import os
import sys
import glob
import base64
import io
import urllib.parse
import webbrowser
from datetime import datetime
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

import unicodedata
import re

PORT = 8080
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

def normalize_key(s):
    if not s: return ''
    s = unicodedata.normalize('NFD', str(s)).encode('ascii', 'ignore').decode('utf-8')
    return re.sub(r'[^a-z0-9]', '', s.lower())

# Comprehensive City -> (Country, Region, Province) Mapping
CITY_LOOKUP = {
    # Brasile
    'natal': ('BRASILE', 'RIO GRANDE DO NORTE', 'NATAL'),
    'pipa': ('BRASILE', 'RIO GRANDE DO NORTE', 'TIBAU DO SUL'),
    'tibaudosul': ('BRASILE', 'RIO GRANDE DO NORTE', 'TIBAU DO SUL'),
    'belohorizonte': ('BRASILE', 'MINAS GERAIS', 'BELO HORIZONTE'),
    'sanpaolo': ('BRASILE', 'SAO PAULO', 'SAO PAULO'),
    'saopaulo': ('BRASILE', 'SAO PAULO', 'SAO PAULO'),
    'portoalegre': ('BRASILE', 'RIO GRANDE DO SUL', 'PORTO ALEGRE'),
    'riodejaneiro': ('BRASILE', 'RIO DE JANEIRO', 'RIO DE JANEIRO'),
    'brasilia': ('BRASILE', 'DISTRITO FEDERAL', 'BRASILIA'),
    'salvador': ('BRASILE', 'BAHIA', 'SALVADOR'),
    'fortaleza': ('BRASILE', 'CEARA', 'FORTALEZA'),
    'recife': ('BRASILE', 'PERNAMBUCO', 'RECIFE'),
    'curitiba': ('BRASILE', 'PARANA', 'CURITIBA'),
    'manaus': ('BRASILE', 'AMAZONAS', 'MANAUS'),
    'florianopolis': ('BRASILE', 'SANTA CATARINA', 'FLORIANOPOLIS'),
    'goiania': ('BRASILE', 'GOIAS', 'GOIANIA'),
    'belem': ('BRASILE', 'PARA', 'BELEM'),
    'joaopessoa': ('BRASILE', 'PARAIBA', 'JOAO PESSOA'),
    'maceio': ('BRASILE', 'ALAGOAS', 'MACEIO'),
    'aracaju': ('BRASILE', 'SERGIPE', 'ARACAJU'),
    'teresina': ('BRASILE', 'PIAUI', 'TERESINA'),
    'saoluis': ('BRASILE', 'MARANHAO', 'SAO LUIS'),
    'campinas': ('BRASILE', 'SAO PAULO', 'CAMPINAS'),
    'santos': ('BRASILE', 'SAO PAULO', 'SANTOS'),
    'ribeiraopreto': ('BRASILE', 'SAO PAULO', 'RIBEIRAO PRETO'),
    
    # Italia
    'roma': ('ITALIA', 'LAZIO', 'ROMA'),
    'milano': ('ITALIA', 'LOMBARDIA', 'MILANO'),
    'napoli': ('ITALIA', 'CAMPANIA', 'NAPOLI'),
    'torino': ('ITALIA', 'PIEMONTE', 'TORINO'),
    'palermo': ('ITALIA', 'SICILIA', 'PALERMO'),
    'genova': ('ITALIA', 'LIGURIA', 'GENOVA'),
    'bologna': ('ITALIA', 'EMILIA ROMAGNA', 'BOLOGNA'),
    'firenze': ('ITALIA', 'TOSCANA', 'FIRENZE'),
    'bari': ('ITALIA', 'PUGLIA', 'BARI'),
    'catania': ('ITALIA', 'SICILIA', 'CATANIA'),
    'venezia': ('ITALIA', 'VENETO', 'VENEZIA'),
    'verona': ('ITALIA', 'VENETO', 'VERONA'),
    'messina': ('ITALIA', 'SICILIA', 'MESSINA'),
    'padova': ('ITALIA', 'VENETO', 'PADOVA'),
    'trieste': ('ITALIA', 'FRIULI VENEZIA GIULIA', 'TRIESTE'),
    'taranto': ('ITALIA', 'PUGLIA', 'TARANTO'),
    'brescia': ('ITALIA', 'LOMBARDIA', 'BRESCIA'),
    'parma': ('ITALIA', 'EMILIA ROMAGNA', 'PARMA'),
    'prato': ('ITALIA', 'TOSCANA', 'PRATO'),
    'modena': ('ITALIA', 'EMILIA ROMAGNA', 'MODENA'),
    'reggiocalabria': ('ITALIA', 'CALABRIA', 'REGGIO CALABRIA'),
    'reggioemilia': ('ITALIA', 'EMILIA ROMAGNA', 'REGGIO EMILIA'),
    'perugia': ('ITALIA', 'UMBRIA', 'PERUGIA'),
    'ravenna': ('ITALIA', 'EMILIA ROMAGNA', 'RAVENNA'),
    'livorno': ('ITALIA', 'TOSCANA', 'LIVORNO'),
    'cagliari': ('ITALIA', 'SARDEGNA', 'CAGLIARI'),
    'foggia': ('ITALIA', 'PUGLIA', 'FOGGIA'),
    'rimini': ('ITALIA', 'EMILIA ROMAGNA', 'RIMINI'),
    'salerno': ('ITALIA', 'CAMPANIA', 'SALERNO'),
    'ferrara': ('ITALIA', 'EMILIA ROMAGNA', 'FERRARA'),
    'sassari': ('ITALIA', 'SARDEGNA', 'SASSARI'),
    'latina': ('ITALIA', 'LAZIO', 'LATINA'),
    'monza': ('ITALIA', 'LOMBARDIA', 'MONZA E DELLA BRIANZA'),
    'bergamo': ('ITALIA', 'LOMBARDIA', 'BERGAMO'),
    'pescara': ('ITALIA', 'ABRUZZO', 'PESCARA'),
    'siracusa': ('ITALIA', 'SICILIA', 'SIRACUSA'),
    'forli': ('ITALIA', 'EMILIA ROMAGNA', 'FORLI-CESENA'),
    'trento': ('ITALIA', 'TRENTINO-ALTO ADIGE', 'TRENTO'),
    'vicenza': ('ITALIA', 'VENETO', 'VICENZA'),
    'terni': ('ITALIA', 'UMBRIA', 'TERNI'),
    'bolzano': ('ITALIA', 'TRENTINO-ALTO ADIGE', 'BOLZANO'),
    'novara': ('ITALIA', 'PIEMONTE', 'NOVARA'),
    'piacenza': ('ITALIA', 'EMILIA ROMAGNA', 'PIACENZA'),
    'ancona': ('ITALIA', 'MARCHE', 'ANCONA'),
    'udine': ('ITALIA', 'FRIULI VENEZIA GIULIA', 'UDINE'),
    'arezzo': ('ITALIA', 'TOSCANA', 'AREZZO'),
    'cesena': ('ITALIA', 'EMILIA ROMAGNA', 'FORLI-CESENA'),
    'lecce': ('ITALIA', 'PUGLIA', 'LECCE'),
    'pesaro': ('ITALIA', 'MARCHE', 'PESARO E URBINO'),
    'alessandria': ('ITALIA', 'PIEMONTE', 'ALESSANDRIA'),
    'laspezia': ('ITALIA', 'LIGURIA', 'LA SPEZIA'),
    'pisa': ('ITALIA', 'TOSCANA', 'PISA'),
    'pistoia': ('ITALIA', 'TOSCANA', 'PISTOIA'),
    'lucca': ('ITALIA', 'TOSCANA', 'LUCCA'),
    'trapani': ('ITALIA', 'SICILIA', 'TRAPANI'),
    
    # Europa & Mondo
    'lima': ('PERU', 'LIMA', 'LIMA'),
    'cusco': ('PERU', 'CUSCO', 'CUSCO'),
    'parigi': ('FRANCIA', 'ILE-DE-FRANCE', 'PARIGI'),
    'paris': ('FRANCIA', 'ILE-DE-FRANCE', 'PARIGI'),
    'lione': ('FRANCIA', 'AUVERGNE-RHONE-ALPES', 'LIONE'),
    'marsiglia': ('FRANCIA', 'PROVENCE-ALPES-COTE D AZUR', 'MARSIGLIA'),
    'nizza': ('FRANCIA', 'PROVENCE-ALPES-COTE D AZUR', 'NIZZA'),
    'bordeaux': ('FRANCIA', 'NOUVELLE-AQUITAINE', 'BORDEAUX'),
    'londra': ('REGNO UNITO', 'ENGLAND', 'GREATER LONDON'),
    'london': ('REGNO UNITO', 'ENGLAND', 'GREATER LONDON'),
    'manchester': ('REGNO UNITO', 'ENGLAND', 'MANCHESTER'),
    'birmingham': ('REGNO UNITO', 'ENGLAND', 'BIRMINGHAM'),
    'edimburgo': ('REGNO UNITO', 'SCOTLAND', 'EDINBURGH'),
    'zurigo': ('SVIZZERA', 'ZURIGO', 'ZURIGO'),
    'zurich': ('SVIZZERA', 'ZURIGO', 'ZURIGO'),
    'ginevra': ('SVIZZERA', 'GINEVRA', 'GINEVRA'),
    'lugano': ('SVIZZERA', 'TICINO', 'LUGANO'),
    'berna': ('SVIZZERA', 'BERNA', 'BERNA'),
    'basilea': ('SVIZZERA', 'BASILEA', 'BASILEA'),
    'santiago': ('CILE', 'SANTIAGO', 'SANTIAGO'),
    'valparaiso': ('CILE', 'VALPARAISO', 'VALPARAISO'),
    'buenosaires': ('ARGENTINA', 'BUENOS AIRES', 'BUENOS AIRES'),
    'cordoba': ('ARGENTINA', 'CORDOBA', 'CORDOBA'),
    'rosario': ('ARGENTINA', 'SANTA FE', 'ROSARIO'),
    'mendoza': ('ARGENTINA', 'MENDOZA', 'MENDOZA'),
    'madrid': ('SPAGNA', 'MADRID', 'MADRID'),
    'barcellona': ('SPAGNA', 'CATALONIA', 'BARCELLONA'),
    'barcelona': ('SPAGNA', 'CATALONIA', 'BARCELLONA'),
    'valencia': ('SPAGNA', 'VALENCIA', 'VALENCIA'),
    'siviglia': ('SPAGNA', 'ANDALUCIA', 'SIVIGLIA'),
    'lisbona': ('PORTOGALLO', 'LISBONA', 'LISBONA'),
    'porto': ('PORTOGALLO', 'PORTO', 'PORTO'),
    'berlino': ('GERMANIA', 'BERLINO', 'BERLINO'),
    'berlin': ('GERMANIA', 'BERLINO', 'BERLINO'),
    'monaco': ('GERMANIA', 'BAVIERA', 'MONACO'),
    'monacodibaviera': ('GERMANIA', 'BAVIERA', 'MONACO'),
    'francoforte': ('GERMANIA', 'ASSIA', 'FRANCOFORTE'),
    'amburgo': ('GERMANIA', 'AMBURGO', 'AMBURGO'),
    'vienna': ('AUSTRIA', 'VIENNA', 'VIENNA'),
    'bruxelles': ('BELGIO', 'BRUXELLES', 'BRUXELLES'),
    'amsterdam': ('PAESI BASSI', 'OLANDA SETTENTRIONALE', 'AMSTERDAM'),
    'newyork': ('STATI UNITI', 'NEW YORK', 'NEW YORK'),
    'losangeles': ('STATI UNITI', 'CALIFORNIA', 'LOS ANGELES'),
    'sanfrancisco': ('STATI UNITI', 'CALIFORNIA', 'SAN FRANCISCO'),
    'miami': ('STATI UNITI', 'FLORIDA', 'MIAMI'),
    'orlando': ('STATI UNITI', 'FLORIDA', 'ORLANDO'),
    'chicago': ('STATI UNITI', 'ILLINOIS', 'CHICAGO'),
    'toronto': ('CANADA', 'ONTARIO', 'TORONTO'),
    'montreal': ('CANADA', 'QUEBEC', 'MONTREAL'),
    'vancouver': ('CANADA', 'BRITISH COLUMBIA', 'VANCOUVER'),
    'cittadelmessico': ('MESSICO', 'CIUDAD DE MEXICO', 'CIUDAD DE MEXICO'),
    'bogota': ('COLOMBIA', 'BOGOTA', 'BOGOTA'),
    'montevideo': ('URUGUAY', 'MONTEVIDEO', 'MONTEVIDEO'),
    'asuncion': ('PARAGUAY', 'ASUNCION', 'ASUNCION'),
    'tokyo': ('GIAPPONE', 'TOKYO', 'TOKYO'),
    'pechino': ('CINA', 'BEIJING', 'PECHINO'),
    'shanghai': ('CINA', 'SHANGHAI', 'SHANGHAI'),
    'sydney': ('AUSTRALIA', 'NEW SOUTH WALES', 'SYDNEY'),
    'dubai': ('EMIRATI ARABI UNITI', 'DUBAI', 'DUBAI')
}

def get_desktop_excel_path():
    user_home = os.path.expanduser("~")
    desktop_dir = os.path.join(user_home, "Desktop")
    
    for name in ["Mappe.xlsx", "mappe.xlsx", "Mappe.xls", "mappe.xls"]:
        p = os.path.join(desktop_dir, name)
        if os.path.exists(p):
            return p
            
    matches = glob.glob(os.path.join(desktop_dir, "*mappe*.xlsx")) + glob.glob(os.path.join(desktop_dir, "*mappe*.xls"))
    if matches:
        return matches[0]
        
    for name in ["Mappe.xlsx", "mappe.xlsx"]:
        p = os.path.join(BASE_DIR, name)
        if os.path.exists(p):
            return p
            
    return None

def enrich_and_save_excel(file_path, sheet_name=None):
    """
    Checks the Excel file for any rows where CITTA is present but PAESE, REGIONE, or PROVINCIA are empty.
    Fills in the missing values from CITY_LOOKUP and saves the Excel file directly to disk.
    """
    if not file_path or not os.path.exists(file_path):
        return False
    try:
        wb = openpyxl.load_workbook(file_path)
        sheets_to_check = [sheet_name] if (sheet_name and sheet_name in wb.sheetnames) else wb.sheetnames
        
        country_keywords = ["paese", "stato", "country", "nazione", "state", "nazioni", "stati"]
        region_keywords = ["regione", "region", "regioni", "reg", "distretto", "departement", "departamento", "estado"]
        province_keywords = ["provincia", "prov", "province", "sigla", "county"]
        city_keywords = ["citta", "città", "city", "cidade", "municipio", "comune"]
        
        total_modified = 0
        for sname in sheets_to_check:
            ws = wb[sname]
            header_row_idx = None
            col_c, col_r, col_p, col_city = None, None, None, None
            
            for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                cells = [str(c).strip().lower() for c in row if c is not None and str(c).strip()]
                if any(any(kw == c or kw in c for kw in city_keywords) for c in cells):
                    header_row_idx = r_idx
                    break
            
            if not header_row_idx:
                continue
                
            for c_idx in range(1, ws.max_column + 1):
                v = ws.cell(row=header_row_idx, column=c_idx).value
                if not v:
                    continue
                v_str = str(v).strip().lower()
                if not col_c and any(kw in v_str for kw in country_keywords):
                    col_c = c_idx
                elif not col_r and any(kw in v_str for kw in region_keywords):
                    col_r = c_idx
                elif not col_p and any(kw in v_str for kw in province_keywords):
                    col_p = c_idx
                elif not col_city and any(kw in v_str for kw in city_keywords):
                    col_city = c_idx
                    
            if not col_city:
                continue
                
            for r_idx in range(header_row_idx + 1, ws.max_row + 1):
                city_val = ws.cell(row=r_idx, column=col_city).value
                if not city_val:
                    continue
                c_norm = normalize_key(city_val)
                geo = CITY_LOOKUP.get(c_norm)
                if geo:
                    match_country, match_region, match_province = geo
                    if col_c:
                        c_cell = ws.cell(row=r_idx, column=col_c)
                        if c_cell.value is None or str(c_cell.value).strip() == "":
                            c_cell.value = match_country
                            total_modified += 1
                    if col_r:
                        r_cell = ws.cell(row=r_idx, column=col_r)
                        if r_cell.value is None or str(r_cell.value).strip() == "":
                            r_cell.value = match_region
                            total_modified += 1
                    if col_p:
                        p_cell = ws.cell(row=r_idx, column=col_p)
                        if p_cell.value is None or str(p_cell.value).strip() == "":
                            p_cell.value = match_province
                            total_modified += 1
                            
        if total_modified > 0:
            wb.save(file_path)
            print(f"Aggiornate {total_modified} celle geografiche e salvato il file Excel: {file_path}")
            return True
    except PermissionError:
        print(f"File Excel aperto da un'altra applicazione, aggiornamento su disco posticipato: {file_path}")
    except Exception as e:
        print(f"Errore arricchimento Excel su disco: {e}")
    return False

def parse_excel_file(file_path, sheet_name=None):
    if not file_path or not os.path.exists(file_path):
        return {
            "success": False,
            "error": f"File non trovato sul Desktop: {file_path or 'Mappe.xlsx'}"
        }

    # Automatically enrich missing Paese, Regione, Provincia from Citta and save to Excel file
    enrich_and_save_excel(file_path, sheet_name)

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheets = wb.sheetnames
        
        if sheet_name and sheet_name in sheets:
            ws = wb[sheet_name]
            selected_sheet = sheet_name
        else:
            ws = wb.active
            selected_sheet = ws.title

        rows_data = list(ws.iter_rows(values_only=True))
        if not rows_data:
            return {
                "success": False,
                "error": "Il foglio Excel selezionato è vuoto."
            }

        country_keywords = ["paese", "stato", "country", "nazione", "state", "nazioni", "stati"]
        region_keywords = ["regione", "region", "regioni", "reg", "distretto", "departement", "departamento", "estado"]
        province_keywords = ["provincia", "prov", "province", "sigla", "county"]
        city_keywords = ["citta", "città", "city", "cidade", "municipio", "comune"]
        channel_keywords = ["canali", "canale", "prenotazione", "prenotazioni", "preovenienza", "provenienza", "booking_channel", "channel", "fonte", "tipo", "tipologia", "agenzia", "agency", "origine"]
        hotel_keywords = ["hotel", "struttura", "albergo", "resort", "strutture", "unita", "unità", "proprieta", "proprietà", "stabilimento"]
        year_keywords = ["anno", "year", "exercicio", "esercizio"]
        month_keywords = ["mese", "month", "mes"]
        metric_keywords = ["clienti", "ospiti", "presenze", "pernottamenti", "vendite", "vendita", "valore", "importo", "fatturato", "sales", "revenue", "qta", "quantita", "quantità", "valori", "totale", "total", "%", "percentuale", "quota", "share"]

        best_score = -1
        best_row_idx = 0
        
        for idx, row in enumerate(rows_data):
            if not row:
                continue
            cells = [str(c).strip().lower() for c in row if c is not None and str(c).strip()]
            if len(cells) < 1:
                continue
                
            score = 0
            for c in cells:
                if any(kw == c or kw in c for kw in country_keywords):
                    score += 10
                if any(kw == c or kw in c for kw in region_keywords):
                    score += 10
                if any(kw == c or kw in c for kw in city_keywords):
                    score += 10
                if any(kw == c or kw in c for kw in channel_keywords):
                    score += 8
                if any(kw == c or kw in c for kw in province_keywords):
                    score += 10
                if any(kw == c or kw in c for kw in metric_keywords):
                    score += 8
                if len(c) > 1 and not c.replace('.', '', 1).replace(',', '', 1).isdigit():
                    score += 1
                    
            if score > best_score:
                best_score = score
                best_row_idx = idx

        header_idx = best_row_idx
        raw_headers = rows_data[header_idx]
        
        headers = []
        for i, h in enumerate(raw_headers):
            if h is not None and str(h).strip():
                headers.append((i, str(h).strip()))
            else:
                has_data = any(rows_data[r][i] is not None for r in range(header_idx + 1, min(len(rows_data), header_idx + 50)) if i < len(rows_data[r]))
                if has_data:
                    headers.append((i, f"Colonna_{i+1}"))

        if not headers:
            return {
                "success": False,
                "error": "Nessuna colonna valida identificata nel foglio."
            }

        records = []
        for r_idx in range(header_idx + 1, len(rows_data)):
            row = rows_data[r_idx]
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue
                
            item = {}
            has_val = False
            for col_idx, col_name in headers:
                if col_idx < len(row):
                    val = row[col_idx]
                    if val is not None:
                        if isinstance(val, str):
                            val = val.strip()
                        item[col_name] = val
                        has_val = True
                    else:
                        item[col_name] = None
                else:
                    item[col_name] = None
            if has_val:
                records.append(item)

        detected = {
            "hotel": None,
            "year": None,
            "month": None,
            "country": None,
            "region": None,
            "province": None,
            "city": None,
            "channel": None,
            "metrics": [],
            "all": [h[1] for h in headers]
        }

        for _, name in headers:
            nl = name.lower()
            if not detected["hotel"] and any(kw == nl or kw in nl for kw in hotel_keywords):
                detected["hotel"] = name
            elif not detected["year"] and any(kw == nl or kw in nl for kw in year_keywords):
                detected["year"] = name
            elif not detected["month"] and any(kw == nl or kw in nl for kw in month_keywords):
                detected["month"] = name
            elif not detected["country"] and any(kw in nl for kw in country_keywords):
                detected["country"] = name
            elif not detected["region"] and any(kw in nl for kw in region_keywords):
                detected["region"] = name
            elif not detected["city"] and any(kw in nl for kw in city_keywords):
                detected["city"] = name
            elif not detected["channel"] and any(kw in nl for kw in channel_keywords):
                detected["channel"] = name
            elif not detected["province"] and any(kw in nl for kw in province_keywords):
                detected["province"] = name

        # Automatic Geographic Enrichment: When CITTA is present, auto-fill missing Country, Region, Province
        c_col = detected["country"] or "PAESE"
        r_col = detected["region"] or "REGIONE"
        p_col = detected["province"] or "PROVINCIA"
        city_col = detected["city"]

        for r in records:
            city_val = r.get(city_col) if city_col else None
            if city_val:
                c_norm = normalize_key(city_val)
                geo_match = CITY_LOOKUP.get(c_norm)
                if geo_match:
                    match_country, match_region, match_province = geo_match
                    if not r.get(c_col) or str(r.get(c_col)).strip() == "":
                        r[c_col] = match_country
                        if not detected["country"]: detected["country"] = c_col
                    if not r.get(r_col) or str(r.get(r_col)).strip() == "":
                        r[r_col] = match_region
                        if not detected["region"]: detected["region"] = r_col
                    if not r.get(p_col) or str(r.get(p_col)).strip() == "":
                        r[p_col] = match_province
                        if not detected["province"]: detected["province"] = p_col

        # Extract distinct hotels if hotel column is detected
        hotels_found = []
        if detected["hotel"]:
            hotels_found = sorted(list(set(str(r[detected["hotel"]]).strip() for r in records if r.get(detected["hotel"]))))

        # Exclude summary/total rows (e.g. bottom total row)
        dim_cols_check = [detected["hotel"], detected["country"], detected["region"], detected["province"]]
        dim_cols_check = [d for d in dim_cols_check if d]
        if dim_cols_check:
            clean_records = []
            for r in records:
                has_dim = any(r.get(d) is not None and str(r.get(d)).strip() != "" for d in dim_cols_check)
                is_total = any(str(v).strip().lower() in ["totale", "total", "somma", "grand total"] for v in r.values() if v is not None and isinstance(v, str))
                if has_dim and not is_total:
                    clean_records.append(r)
            records = clean_records

        # Classify metric columns (excluding dimensions)
        dim_cols = [detected["hotel"], detected["year"], detected["month"], detected["country"], detected["region"], detected["province"]]
        for _, name in headers:
            if name in dim_cols:
                continue
            num_count = 0
            val_count = 0
            for item in records:
                v = item.get(name)
                if v is not None:
                    val_count += 1
                    if isinstance(v, (int, float)):
                        num_count += 1
                    elif isinstance(v, str):
                        try:
                            float(v.replace(',', '.').replace('%', '').strip())
                            num_count += 1
                        except ValueError:
                            pass
            if val_count > 0 and (num_count / val_count >= 0.4):
                detected["metrics"].append(name)

        if not detected["metrics"] and records:
            for _, name in headers:
                if name not in dim_cols:
                    detected["metrics"].append(name)

        last_mod = os.path.getmtime(file_path)
        last_mod_str = datetime.fromtimestamp(last_mod).strftime("%d/%m/%Y %H:%M:%S")

        return {
            "success": True,
            "filename": os.path.basename(file_path),
            "filepath": file_path,
            "last_modified": last_mod_str,
            "sheets": sheets,
            "current_sheet": selected_sheet,
            "detected_columns": detected,
            "hotels": hotels_found,
            "row_count": len(records),
            "data": records
        }

    except PermissionError:
        return {
            "success": False,
            "error": f"Il file '{os.path.basename(file_path)}' è aperto in Excel. Chiudilo e riprova."
        }
    except Exception as e:
        return {
            "success": False,
            "error": f"Errore durante la lettura di Excel: {str(e)}"
        }

def save_map_to_excel(image_bytes, sheet_name=None, custom_path=None):
    file_path = custom_path or get_desktop_excel_path()
    if not file_path or not os.path.exists(file_path):
        return {
            "success": False,
            "error": f"File Excel non trovato sul Desktop: {file_path or 'Mappe.xlsx'}"
        }

    try:
        temp_img_path = os.path.join(BASE_DIR, "temp_map_export.png")
        pil_img = PILImage.open(io.BytesIO(image_bytes))
        
        # Max width 800px for high clarity in Excel
        max_w = 800
        if pil_img.width > max_w:
            ratio = max_w / float(pil_img.width)
            new_h = int(float(pil_img.height) * ratio)
            pil_img = pil_img.resize((max_w, new_h), PILImage.LANCZOS)
            
        pil_img.save(temp_img_path, format="PNG")

        wb = openpyxl.load_workbook(file_path)
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.active

        # Find bounds of existing data table
        min_r, max_r, min_c, max_c = None, None, None, None
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            for c_idx, val in enumerate(row, start=1):
                if val is not None and str(val).strip():
                    if min_r is None or r_idx < min_r: min_r = r_idx
                    if max_r is None or r_idx > max_r: max_r = r_idx
                    if min_c is None or c_idx < min_c: min_c = c_idx
                    if max_c is None or c_idx > max_c: max_c = c_idx

        # Calculate position 2 columns to the right of the table
        target_col_idx = (max_c or 6) + 2
        target_row_idx = min_r or 1
        anchor_cell = f"{get_column_letter(target_col_idx)}{target_row_idx}"

        # Clear existing images on sheet
        ws._images = []

        # Insert image
        img = OpenpyxlImage(temp_img_path)
        img.anchor = anchor_cell
        ws.add_image(img)

        wb.save(file_path)

        return {
            "success": True,
            "anchor": anchor_cell,
            "filename": os.path.basename(file_path),
            "message": f"Mappa inserita con successo nel file Excel '{os.path.basename(file_path)}' alla posizione {anchor_cell}!"
        }

    except PermissionError:
        return {
            "success": False,
            "error": f"Il file '{os.path.basename(file_path)}' è attualmente aperto in Microsoft Excel. Chiudi il file Excel sul tuo computer e riprova a cliccare sul pulsante."
        }
    except Exception as e:
        return {
            "success": False,
            "error": f"Errore durante il salvataggio in Excel: {str(e)}"
        }

class MapDashboardHandler(http.server.SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=BASE_DIR, **kwargs)

    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        if parsed.path == "/api/data":
            params = urllib.parse.parse_qs(parsed.query)
            sheet = params.get("sheet", [None])[0]
            custom_path = params.get("path", [None])[0]
            
            excel_path = custom_path or get_desktop_excel_path()
            res = parse_excel_file(excel_path, sheet)
            
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Access-Control-Allow-Origin", "*")
            self.send_header("Cache-Control", "no-cache, no-store, must-revalidate")
            self.end_headers()
            self.wfile.write(json.dumps(res, ensure_ascii=False).encode("utf-8"))
            return

        if parsed.path == "/api/status":
            excel_path = get_desktop_excel_path()
            status = {
                "running": True,
                "excel_path": excel_path,
                "excel_exists": excel_path is not None and os.path.exists(excel_path)
            }
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Access-Control-Allow-Origin", "*")
            self.end_headers()
            self.wfile.write(json.dumps(status).encode("utf-8"))
            return

        super().do_GET()

    def do_POST(self):
        parsed = urllib.parse.urlparse(self.path)
        if parsed.path == "/api/export-excel":
            content_length = int(self.headers.get('Content-Length', 0))
            body = self.rfile.read(content_length)
            
            try:
                data = json.loads(body.decode('utf-8'))
                img_data_url = data.get("image_base64", "")
                sheet = data.get("sheet", None)
                custom_path = data.get("path", None)

                if "," in img_data_url:
                    img_data_url = img_data_url.split(",", 1)[1]
                
                img_bytes = base64.b64decode(img_data_url)
                res = save_map_to_excel(img_bytes, sheet, custom_path)
                
                self.send_response(200)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.send_header("Access-Control-Allow-Origin", "*")
                self.end_headers()
                self.wfile.write(json.dumps(res, ensure_ascii=False).encode("utf-8"))
            except Exception as e:
                self.send_response(200)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.send_header("Access-Control-Allow-Origin", "*")
                self.end_headers()
                self.wfile.write(json.dumps({"success": False, "error": str(e)}).encode("utf-8"))
            return

        super().do_POST()

def open_browser_delayed(url):
    import time
    import subprocess
    time.sleep(0.4)
    try:
        edge_path = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
        chrome_path = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
        if os.path.exists(edge_path):
            subprocess.Popen([edge_path, url])
        elif os.path.exists(chrome_path):
            subprocess.Popen([chrome_path, url])
        else:
            webbrowser.open(url)
    except Exception as e:
        try:
            webbrowser.open(url)
        except Exception:
            pass

def run_server(port=PORT):
    http.server.ThreadingHTTPServer.allow_reuse_address = True
    actual_port = port
    server = None
    for p in range(port, port + 10):
        try:
            server = http.server.ThreadingHTTPServer(("127.0.0.1", p), MapDashboardHandler)
            actual_port = p
            break
        except OSError:
            continue
            
    if not server:
        print(f"Impossibile avviare il server sulle porte {port}-{port+10}")
        return

    url = f"http://127.0.0.1:{actual_port}"
    print(f"==================================================")
    print(f"   DASHBOARD MAPPE HOTEL AVVIATA CON SUCCESSO!    ")
    print(f"==================================================")
    print(f"Indirizzo Web: {url}")
    excel = get_desktop_excel_path()
    if excel:
        print(f"File Excel rilevato: {excel}")
    else:
        print("In attesa di file 'Mappe.xlsx' sul Desktop...")
    print(f"Apertura automatica del browser in corso...")
    print(f"Premi CTRL+C per chiudere la dashboard.")
    print(f"==================================================")
    
    import threading
    t = threading.Thread(target=open_browser_delayed, args=(url,), daemon=True)
    t.start()
        
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nArresto del server.")
        server.server_close()

if __name__ == "__main__":
    p = int(sys.argv[1]) if len(sys.argv) > 1 else PORT
    run_server(p)
